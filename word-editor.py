from copy import deepcopy
from datetime import datetime

from docx import Document
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from openpyxl import load_workbook

PLIK_WORD = "szablon.docx"
PLIK_EXCEL = "dane.xlsx"
PLIK_WYJSCIOWY = "wynik.docx"

# Znaczniki w dokumencie w kolejności kolumn Excela A, B, C.
# Kolumna D to liczba kopii strony. Pierwszy wiersz Excela to nagłówki.
POLA = ["{IMIE}", "{NAZWISKO}", "{DATA}"]
# Miejsce, w które trafi numer strony, np. 1/2, 2/2
POLE_NUMERU = "{NUMER}"


def tekst(wartosc):
    if isinstance(wartosc, datetime):
        return wartosc.strftime("%d.%m.%Y")
    return "" if wartosc is None else str(wartosc)


def zamien(akapit, stary, nowy):
    # Word potrafi podzielić tekst na kilka fragmentów, więc szukamy w całym akapicie
    wezly = akapit.findall(".//" + qn("w:t"))
    calosc = "".join(w.text or "" for w in wezly)
    if stary not in calosc:
        return
    wezly[0].text = calosc.replace(stary, nowy)
    for w in wezly[1:]:
        w.text = ""


def podzial_strony():
    akapit = OxmlElement("w:p")
    run = OxmlElement("w:r")
    br = OxmlElement("w:br")
    br.set(qn("w:type"), "page")
    run.append(br)
    akapit.append(run)
    return akapit


# Wczytaj wiersze z Excela: 3 wartości + liczba kopii
arkusz = load_workbook(PLIK_EXCEL, data_only=True).active
wiersze = [w[:4] for w in arkusz.iter_rows(min_row=2, values_only=True) if w[0] is not None]

doc = Document(PLIK_WORD)

# Zapamiętaj treść szablonu i usuń ją z dokumentu (ustawienia strony zostają)
body = doc.element.body
ustawienia_strony = body.find(qn("w:sectPr"))
szablon = [el for el in body if el is not ustawienia_strony]
for el in szablon:
    body.remove(el)

pierwsza_strona = True
for *dane, numerator in wiersze:
    numerator = int(numerator)
    wartosci = {pole: tekst(w) for pole, w in zip(POLA, dane)}
    for nr in range(1, numerator + 1):
        if not pierwsza_strona:
            ustawienia_strony.addprevious(podzial_strony())
        pierwsza_strona = False
        wartosci[POLE_NUMERU] = f"{nr}/{numerator}"
        for el in szablon:
            kopia = deepcopy(el)
            for akapit in kopia.iter(qn("w:p")):
                for pole, wartosc in wartosci.items():
                    zamien(akapit, pole, wartosc)
            ustawienia_strony.addprevious(kopia)

doc.save(PLIK_WYJSCIOWY)
print(f"Zapisano: {PLIK_WYJSCIOWY} ({len(wiersze)} wierszy z Excela)")
